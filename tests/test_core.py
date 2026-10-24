from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from app.api import app
from app.config import visible_roles
from app.loaders import load_docx, load_text
from app.retrieval import rrf
from app.store import bm25_tokens, chunk_document


def test_rrf_prefers_items_ranked_high_in_both_lists():
    fused = rrf([["a", "b", "c"], ["b", "a", "d"]])
    assert [cid for cid, _ in fused][:2] in (["a", "b"], ["b", "a"])
    assert fused[-1][0] in ("c", "d")


def test_visible_roles():
    assert visible_roles("admin") is None
    assert visible_roles("finance") == ["finance", "all"]


def test_bm25_tokens_drop_stopwords_and_keep_numbers():
    assert bm25_tokens("What is the 10-K for Apple?") == ["10-k", "apple"]


def test_docx_loader_splits_on_headings(tmp_path: Path):
    p = tmp_path / "policy.docx"
    d = Document()
    d.add_heading("Leave", level=1)
    d.add_paragraph("Employees get 24 days.")
    d.add_heading("Sick", level=1)
    d.add_paragraph("Twelve days of sick leave.")
    d.save(p)
    sections = load_docx(p)
    assert [s.section for s in sections] == ["Leave", "Sick"]


def test_chunks_keep_metadata_and_respect_size(tmp_path: Path, monkeypatch):
    import app.store as store

    monkeypatch.setattr(store, "ROOT", tmp_path)
    p = tmp_path / "doc.md"
    p.write_text("# Intro\n" + ("Sentence number one is here. " * 400), encoding="utf-8")
    chunks = chunk_document({"path": "doc.md", "title": "Doc", "source": "upload", "allowed_roles": ["hr"]})
    assert len(chunks) > 1
    assert all(c["allowed_roles"] == ["hr"] and c["section"] == "Intro" for c in chunks)
    assert load_text(p)[0].section == "Intro"


def test_auth_rejects_unknown_token_and_non_admin_upload():
    client = TestClient(app)
    assert client.get("/me").status_code == 401
    assert client.get("/me", headers={"Authorization": "Bearer finance-token"}).json() == {"role": "finance"}
    r = client.get("/admin/docs", headers={"Authorization": "Bearer finance-token"})
    assert r.status_code == 403
