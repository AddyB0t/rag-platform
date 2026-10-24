"""Download the public corpus and write data/manifest.json.

Sources and the roles allowed to see them:
  finance  -> SEC EDGAR 10-K filings (HTML)
  research -> arXiv papers (PDF)
  hr       -> internal HR policies (DOCX, synthetic, generated here)
  all      -> Wikipedia articles (HTML)
"""

import json
import sys
import time
from pathlib import Path

import httpx
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from app.config import ROOT, settings  # noqa: E402
from scripts.hr_policies import COMPANY, POLICIES  # noqa: E402

UA = {"User-Agent": "Aaditya Barak research project aadityabarak472@gmail.com"}

SEC_COMPANIES = {"Apple": 320193, "Microsoft": 789019, "NVIDIA": 1045810}

ARXIV = {
    "2005.11401": "Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks",
    "1706.03762": "Attention Is All You Need",
    "2004.04906": "Dense Passage Retrieval for Open-Domain Question Answering",
    "2004.12832": "ColBERT: Efficient and Effective Passage Search",
    "2210.03629": "ReAct: Synergizing Reasoning and Acting in Language Models",
    "2309.15217": "RAGAS: Automated Evaluation of Retrieval Augmented Generation",
    "2104.08663": "BEIR: A Heterogeneous Benchmark for Zero-shot Evaluation of IR Models",
}

WIKIPEDIA = [
    "Retrieval-augmented_generation",
    "Large_language_model",
    "Okapi_BM25",
    "Vector_database",
    "Information_retrieval",
    "Transformer_(deep_learning_architecture)",
    "Apple_Inc.",
    "Microsoft",
    "Nvidia",
    "U.S._Securities_and_Exchange_Commission",
    "Form_10-K",
    "Graphics_processing_unit",
]


def fetch(client: httpx.Client, url: str) -> bytes:
    for attempt in range(4):
        r = client.get(url)
        if r.status_code == 200:
            return r.content
        time.sleep(2 * (attempt + 1))
    raise RuntimeError(f"GET {url} -> {r.status_code}")


def collect_sec(client, out: Path) -> list[dict]:
    docs = []
    for name, cik in SEC_COMPANIES.items():
        sub = json.loads(fetch(client, f"https://data.sec.gov/submissions/CIK{cik:010d}.json"))
        recent = sub["filings"]["recent"]
        i = recent["form"].index("10-K")
        acc, doc, date = recent["accessionNumber"][i], recent["primaryDocument"][i], recent["filingDate"][i]
        url = f"https://www.sec.gov/Archives/edgar/data/{cik}/{acc.replace('-', '')}/{doc}"
        path = out / f"{name.lower()}_10k_{date}.html"
        if not path.exists():
            path.write_bytes(fetch(client, url))
            time.sleep(0.5)
        docs.append({
            "path": str(path.relative_to(ROOT)), "source": "sec", "url": url,
            "title": f"{name} Form 10-K (filed {date})", "allowed_roles": ["finance"],
        })
    return docs


def collect_arxiv(client, out: Path) -> list[dict]:
    docs = []
    for aid, title in ARXIV.items():
        path = out / f"arxiv_{aid}.pdf"
        if not path.exists():
            path.write_bytes(fetch(client, f"https://arxiv.org/pdf/{aid}"))
            time.sleep(3)
        docs.append({
            "path": str(path.relative_to(ROOT)), "source": "arxiv",
            "url": f"https://arxiv.org/abs/{aid}", "title": title, "allowed_roles": ["research"],
        })
    return docs


def collect_wikipedia(client, out: Path) -> list[dict]:
    docs = []
    for page in WIKIPEDIA:
        path = out / f"wiki_{page.replace('(', '').replace(')', '').replace('.', '')}.html"
        if not path.exists():
            path.write_bytes(fetch(client, f"https://en.wikipedia.org/api/rest_v1/page/html/{page}"))
            time.sleep(0.5)
        docs.append({
            "path": str(path.relative_to(ROOT)), "source": "wikipedia",
            "url": f"https://en.wikipedia.org/wiki/{page}",
            "title": page.replace("_", " "), "allowed_roles": ["all"],
        })
    return docs


def build_hr_docx(out: Path) -> list[dict]:
    docs = []
    for title, sections in POLICIES.items():
        path = out / (title.lower().replace(" ", "_").replace("-", "_") + ".docx")
        d = Document()
        d.add_heading(f"{COMPANY} - {title}", level=0)
        for heading, body in sections.items():
            d.add_heading(heading, level=1)
            d.add_paragraph(body)
        d.save(path)
        docs.append({
            "path": str(path.relative_to(ROOT)), "source": "hr", "url": "",
            "title": title, "allowed_roles": ["hr"],
        })
    return docs


def main():
    raw = settings.raw_dir
    for sub in ("sec", "arxiv", "wikipedia", "hr"):
        (raw / sub).mkdir(parents=True, exist_ok=True)
    with httpx.Client(headers=UA, timeout=60, follow_redirects=True) as client:
        docs = (
            collect_sec(client, raw / "sec")
            + collect_arxiv(client, raw / "arxiv")
            + collect_wikipedia(client, raw / "wikipedia")
            + build_hr_docx(raw / "hr")
        )
    manifest = ROOT / "data" / "manifest.json"
    manifest.write_text(json.dumps(docs, indent=2))
    print(f"{len(docs)} documents -> {manifest}")


if __name__ == "__main__":
    main()
